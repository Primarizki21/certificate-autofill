"""Generate docs/report/gt_review.xlsx for manual GT review.

Sheets:
  1. Perubahan v8  - 3 audited tingkat corrections already applied to v8 CSV.
  2. Review Tingkat - tingkat labels with weak/no scale evidence in text.
  3. Review Tanggal - dates empty in GT but a date appears in text.
  4. Review Low-Found - certs where few GT values are found in extracted text.

User fills the 'keputusan' column; decisions get applied to
Ground_Truth_Sertifikat_v8.csv afterwards.
"""

import csv
import os

import openpyxl
from openpyxl.styles import Font, PatternFill

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
RAW_CSV = os.path.join(REPO, "Ground_Truth_Sertifikat.csv")
V8_CSV = os.path.join(REPO, "Ground_Truth_Sertifikat_v8.csv")
OUT = os.path.join(REPO, "docs", "report", "gt_review.xlsx")


def read_rows(path):
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def main():
    raw = {r["Nama File"]: r for r in read_rows(RAW_CSV)}
    v8 = {r["Nama File"]: r for r in read_rows(V8_CSV)}

    wb = openpyxl.Workbook()
    hdr_fill = PatternFill("solid", fgColor="D5E8F0")
    bold = Font(bold=True)
    warn = PatternFill("solid", fgColor="FFF2CC")

    def sheet(name, headers):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for c in ws[1]:
            c.fill = hdr_fill
            c.font = bold
        return ws

    # --- Sheet 1: Perubahan v8 ---
    ws = sheet("Perubahan v8", ["file", "field", "raw", "v8", "bukti", "keputusan"])
    changes = [
        ("1966887_221065_skp.pdf", "Tingkat", "Lainnya", "Nasional",
         "AIESEC in Indonesia (chapter nasional) - user decision"),
        ("1981676_219642_skp.pdf", "Tingkat", "Fakultas", "Nasional",
         "teks literal 'TINGKAT NASIONAL'; template sama dengan 2030372"),
        ("2954283_219642_skp.pdf", "Tingkat", "Fakultas", "Internasional",
         "Institut francais d'Indonesie + Univ. Paris-Saclay + Embassy of France"),
    ]
    for row in changes:
        ws.append(list(row) + [""])

    # --- Sheet 2: Review Tingkat ---
    ws = sheet("Review Tingkat", ["file", "folder", "tingkat_gt", "status", "catatan", "keputusan"])
    cat4 = [
        ("Airno_Faiz.pdf", "Nasional", "CONFIRMED-KEEP",
         "Dataquest bisa diikuti semua mahasiswa nasional (user confirm)"),
        ("FIT_Faiz.pdf", "Internasional", "CONFIRMED-KEEP",
         "ada peserta dari luar negeri, walau tidak ada di teks (user confirm)"),
        ("2439919_221065_skp.pdf", "Nasional", "PENDING", "lomba antar-PTN; teks tanpa bukti skala"),
        ("Gelar Rasa_Muhammad Fazil Irvan Putra.pdf", "Nasional", "PENDING", "lomba antar-PTN; teks tanpa bukti skala"),
        ("hakim_lomba.png", "Nasional", "PENDING", "lomba antar-PTN; teks tanpa bukti skala"),
        ("Hitech_Faiz.pdf", "Nasional", "PENDING", "lomba antar-PTN; teks tanpa bukti skala"),
        ("Primarizki Ahmad Hariyono_data slayer 2_lomba.png", "Nasional", "PENDING", "lomba antar-PTN; teks tanpa bukti skala"),
        ("SSF_Faiz.pdf", "Nasional", "PENDING", "lomba antar-PTN; teks tanpa bukti skala"),
        ("2398703_219642_skp.pdf", "Nasional", "PENDING", "aksi sosial nasional; teks tanpa bukti skala"),
        ("2955331_219642_skp.pdf", "Nasional", "PENDING", "teks tanpa bukti skala"),
        ("1966887_221065_skp.pdf", "Nasional", "PENDING",
         "v8 sudah Nasional; teks tanpa bukti skala (AIESEC chapter - domain)"),
    ]
    for fname, gt, status, note in cat4:
        folder = raw.get(fname, {}).get("Folder", "")
        ws.append([fname, folder, gt, status, note, ""])

    # --- Sheet 3: Review Tanggal ---
    ws = sheet("Review Tanggal", ["file", "folder", "waktu_mulai", "waktu_selesai", "status", "catatan", "keputusan"])
    cat1 = [
        ("Sertif sinem_vene_magang UKM.pdf", "CONFIRMED-KEEP-EMPTY",
         "27 Des 2023 = tanggal tanda tangan, bukan periode kegiatan magang UKM"),
        ("2030325_219642_skp.pdf", "CONFIRMED-KEEP-EMPTY",
         "27 Des 2023 = tanggal tanda tangan, bukan periode kegiatan magang UKM"),
        ("2071996_221065_skp.pdf", "CONFIRMED-KEEP-EMPTY",
         "27 Des 2023 = tanggal tanda tangan, bukan periode kegiatan magang UKM"),
    ]
    for fname, status, note in cat1:
        r = raw.get(fname, {})
        ws.append([fname, r.get("Folder", ""), r.get("Waktu Mulai Pelaksanaan", "-"),
                   r.get("Waktu Selesai Pelaksanaan", "-"), status, note, ""])

    # --- Sheet 4: Review Low-Found ---
    ws = sheet("Review Low-Found", ["file", "folder", "found", "analisis", "rekomendasi", "keputusan"])
    low = [
        ("Girifest_Faiz.pdf", "1/5", "OCR jelek",
         "GT benar; nomor di teks terbaca NOM0R:06/001/E P.GIRLSTAT (B->E, spasi hilang); organisasi HIMASTA garbled -> butuh OCR lebih baik"),
        ("KARSA_Venedict_peserta.pdf", "1/5", "OCR merge + label GT",
         "GT benar secara nilai; tanggal terserap (2023yang diselenggarakanoleh); nama GT bentuk panjang tak literal -> OCR merge"),
        ("Venedict G. P_peserta talkshow.pdf", "1/5", "OCR typo",
         "GT benar; UoinersitasAirlangga (typo), BEMFTMM merged -> OCR"),
        ("2030325_219642_skp.pdf", "1/5", "label GT tak literal",
         "GT 'Magang Universitas Airlangga' vs teks 'Magang UKM Universitas Airlangga'; opsional rapikan label"),
        ("2071996_221065_skp.pdf", "1/5", "label GT tak literal",
         "sama dengan 2030325; opsional rapikan label"),
        ("Sertif sinem_vene_magang UKM.pdf", "1/5", "label GT tak literal",
         "sama dengan 2030325; opsional rapikan label"),
    ]
    for fname, found, cat, rec in low:
        folder = raw.get(fname, {}).get("Folder", "")
        ws.append([fname, folder, found, cat, rec, ""])

    for name in wb.sheetnames:
        ws = wb[name]
        ws.column_dimensions["A"].width = 42
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = max(
                ws.column_dimensions[letter].width or 12, 12
            )

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    wb.save(OUT)
    print(f"wrote {OUT}: sheets={wb.sheetnames}")


if __name__ == "__main__":
    main()
