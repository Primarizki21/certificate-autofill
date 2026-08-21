"""Generate example docx for HYB-COMBINED and HYB-LLM pipelines.

Output:
- docs/report/raw_vs_hyb_combined_examples.docx
- docs/report/raw_vs_hyb_llm_examples.docx

Usage:
  python scripts/generate_hyb_examples.py
"""

import json
import os
import sys

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.path.join(REPO, "docs", "report")
PER_CERT = os.path.join(REPO, "tests", "benchmark_runs", "hyb_llm_20260821_135619", "per_cert_results.json")
RAW_DIR = os.path.join(REPO, "tests", "benchmark_runs", "run_20260728_131835", "extracted_texts")

FIELDS = [
    ("nama_kegiatan_sertifikasi", "Nama Kegiatan"),
    ("waktu_mulai_pelaksanaan", "Tanggal Mulai"),
    ("waktu_selesai_pelaksanaan", "Tanggal Selesai"),
    ("penyelenggara_kegiatan", "Penyelenggara"),
    ("nomor_bukti_fisik_nomor_sertifikasi", "Nomor Sertifikat"),
    ("tingkat", "Tingkat"),
]


def load_per_cert():
    with open(PER_CERT) as f:
        return {r["stem"]: r for r in json.load(f)}


def load_raw():
    out = {}
    for fn in os.listdir(RAW_DIR):
        if fn.endswith(".txt"):
            stem = os.path.splitext(fn)[0]
            with open(os.path.join(RAW_DIR, fn), encoding="utf-8", errors="replace") as f:
                out[stem] = f.read().strip()
    return out


def pick_examples(per_cert, raw, mode="combined"):
    """Pick 5 diverse examples."""
    stems = sorted(per_cert.keys())
    pool = [s for s in stems if s in raw and len(raw[s]) > 50]
    chosen = []
    seen = set()

    def finds(pred):
        for s in pool:
            if s in seen:
                continue
            if pred(s):
                return s
        return None

    if mode == "combined":
        # HYB-COMBINED: no LLM, all offline
        # 1. Router menang (exact)
        s = finds(lambda s: per_cert[s]["router_decision"] != "unrouted"
                  and per_cert[s]["fields"]["tingkat"]["combined_match"] == "exact")
        if s:
            chosen.append(("router", s))
            seen.add(s)
        # 2. Semua field exact
        s = finds(lambda s: all(per_cert[s]["fields"][f]["combined_match"] == "exact"
                                for f in ["nama_kegiatan_sertifikasi", "penyelenggara_kegiatan", "tingkat"]))
        if s:
            chosen.append(("all_exact", s))
            seen.add(s)
        # 3. Nama kegiatan exact (AKT-005 berhasil)
        s = finds(lambda s: per_cert[s]["fields"]["nama_kegiatan_sertifikasi"]["combined_match"] == "exact")
        if s:
            chosen.append(("activity", s))
            seen.add(s)
        # 4. Tingkat wrong (unrouted, no LLM)
        s = finds(lambda s: per_cert[s]["router_decision"] == "unrouted"
                  and per_cert[s]["fields"]["tingkat"]["combined_match"] == "wrong")
        if s:
            chosen.append(("tingkat_wrong", s))
            seen.add(s)
        # 5. Organizer exact
        s = finds(lambda s: per_cert[s]["fields"]["penyelenggara_kegiatan"]["combined_match"] == "exact")
        if s:
            chosen.append(("org", s))
            seen.add(s)
    else:
        # HYB-LLM: with LLM fallback
        # 1. LLM correct (unrouted, tingkat exact)
        s = finds(lambda s: per_cert[s]["router_decision"] == "unrouted"
                  and per_cert[s]["fields"]["tingkat"]["llm_match"] == "exact")
        if s:
            chosen.append(("llm_correct", s))
            seen.add(s)
        # 2. LLM wrong (unrouted, tingkat wrong)
        s = finds(lambda s: per_cert[s]["router_decision"] == "unrouted"
                  and per_cert[s]["fields"]["tingkat"]["llm_match"] == "wrong")
        if s:
            chosen.append(("llm_wrong", s))
            seen.add(s)
        # 3. Router menang (exact)
        s = finds(lambda s: per_cert[s]["router_decision"] != "unrouted"
                  and per_cert[s]["fields"]["tingkat"]["llm_match"] == "exact")
        if s:
            chosen.append(("router", s))
            seen.add(s)
        # 4. Semua field exact
        s = finds(lambda s: all(per_cert[s]["fields"][f]["llm_match"] == "exact"
                                for f in ["nama_kegiatan_sertifikasi", "penyelenggara_kegiatan", "tingkat"]))
        if s:
            chosen.append(("all_exact", s))
            seen.add(s)
        # 5. Nama kegiatan exact
        s = finds(lambda s: per_cert[s]["fields"]["nama_kegiatan_sertifikasi"]["llm_match"] == "exact")
        if s:
            chosen.append(("activity", s))
            seen.add(s)

    # Fill remaining
    for s in pool:
        if len(chosen) >= 5:
            break
        if s not in seen:
            chosen.append(("extra", s))
            seen.add(s)

    order = {"router": 0, "llm_correct": 0, "all_exact": 1, "activity": 2,
             "tingkat_wrong": 3, "llm_wrong": 3, "org": 4, "extra": 5}
    chosen.sort(key=lambda t: (order.get(t[0], 9), t[1]))
    return [s for _, s in chosen[:5]]


def _shade(cell, hexfill):
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hexfill})
    tcPr.append(shd)


def _table_borders(tbl, color="C9D4E4"):
    from docx.oxml.ns import qn
    tblPr = tbl._tbl.tblPr
    borders = tblPr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tblPr.makeelement(qn(f"w:{edge}"), {
            qn("w:val"): "single", qn("w:sz"): "4", qn("w:space"): "0", qn("w:color"): color})
        borders.append(el)
    tblPr.append(borders)


def _raw_box(doc, text):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    p = doc.add_paragraph()
    for line in text.splitlines():
        run = p.add_run(line + "\n")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    pPr = p._p.get_or_add_pPr()
    shd = pPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "F2F6FC"})
    pPr.append(shd)


def render_docx(mode, per_cert, raw, examples, filename, title, subtitle):
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    ACCENT = RGBColor(0x1F, 0x38, 0x64)
    doc = docx.Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    t = doc.add_heading(title, level=0)
    for run in t.runs:
        run.font.color.rgb = ACCENT
    sp = doc.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp.add_run(subtitle).font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    if mode == "combined":
        doc.add_paragraph(
            "Tabel berikut menunjukkan teks MENTAH dari PDF, field hasil HYB-COMBINED "
            "(offline, 0 LLM), dan Ground Truth. Hijau = EXACT, kuning = FUZZY, merah = salah. "
            "HYB-COMBINED menggunakan: router rules (tingkat), AKT-005 (nama kegiatan), "
            "ORG-004 (organizer), PROD-002 (nomor). MACRO 73.7%."
        )
    else:
        doc.add_paragraph(
            "Tabel berikut menunjukkan teks MENTAH dari PDF, field hasil HYB-LLM "
            "(router + LLM fallback), dan Ground Truth. Hijau = EXACT, kuning = FUZZY, merah = salah. "
            "HYB-LLM = HYB-COMBINED + LLM llama3.1:8b untuk 26 cert unrouted. MACRO 74.2%."
        )

    for idx, stem in enumerate(examples, 1):
        cert = per_cert[stem]
        doc.add_heading(f"Contoh {idx} — {stem}", level=1)

        doc.add_heading("Teks mentah hasil ekstraksi PDF:", level=2)
        _raw_box(doc, raw.get(stem, "(tidak ditemukan)"))

        doc.add_heading("Hasil pipeline vs Ground Truth:", level=2)
        tbl = doc.add_table(rows=1 + len(FIELDS), cols=4)
        tbl.style = "Table Grid"
        _table_borders(tbl)

        # Header
        hdr = ["Field", "Pipeline", "GT", "Status"]
        for j, h in enumerate(hdr):
            c = tbl.rows[0].cells[j]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade(c, "1F3864")

        # Data rows
        match_key = "combined_match" if mode == "combined" else "llm_match"
        value_key = "combined" if mode == "combined" else "llm"

        for i, (field, label) in enumerate(FIELDS):
            fv = cert["fields"].get(field, {})
            pipeline_val = fv.get(value_key, "")
            gt_val = fv.get("gt", "")
            status = fv.get(match_key, "wrong")

            row = tbl.rows[i + 1].cells
            for j, val in enumerate([label, pipeline_val or "(kosong)", gt_val or "(kosong)",
                                     {"exact": "EXACT", "fuzzy": "FUZZY", "wrong": "WRONG"}[status]]):
                row[j].text = ""
                r = row[j].paragraphs[0].add_run(val)
                r.font.size = Pt(9)
            _shade(row[3], {"exact": "C6EFCE", "fuzzy": "FFF2CC", "wrong": "FFC7CE"}[status])

        # Router / LLM info
        p = doc.add_paragraph()
        if cert["router_decision"] != "unrouted":
            r = p.add_run(f"Router: {cert['router_decision']} (rule: {cert['router_rule']}) — tanpa LLM")
            r.italic = True
            r.font.size = Pt(10)
        else:
            if mode == "llm":
                llm_answer = fv.get("llm", "")
                llm_match = fv.get("llm_match", "wrong")
                r = p.add_run(f"Router: unrouted → LLM menjawab '{llm_answer}' — {'BENAR' if llm_match == 'exact' else 'SALAH'}")
                r.italic = True
                r.font.size = Pt(10)
            else:
                r = p.add_run("Router: unrouted → tidak ada LLM (offline mode) — tingkat kosong/salah")
                r.italic = True
                r.font.size = Pt(10)

        doc.add_paragraph()

    doc.save(os.path.join(OUT, filename))
    print(f"  {filename}")


def main():
    per_cert = load_per_cert()
    raw = load_raw()

    # HYB-COMBINED examples
    examples_combined = pick_examples(per_cert, raw, mode="combined")
    render_docx("combined", per_cert, raw, examples_combined,
                "raw_vs_hyb_combined_examples.docx",
                "Raw Text vs HYB-COMBINED — Contoh Ekstraksi Nyata",
                "Certificate Autofill Prototype — HYB-COMBINED (offline, 0 LLM)")

    # HYB-LLM examples
    examples_llm = pick_examples(per_cert, raw, mode="llm")
    render_docx("llm", per_cert, raw, examples_llm,
                "raw_vs_hyb_llm_examples.docx",
                "Raw Text vs HYB-LLM — Contoh Ekstraksi Nyata",
                "Certificate Autofill Prototype — HYB-LLM (router + LLM fallback)")


if __name__ == "__main__":
    main()
