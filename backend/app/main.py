import asyncio
import logging
import sys
from contextlib import suppress
from pathlib import Path

_backend_dir = str(Path(__file__).resolve().parent.parent)
if _backend_dir not in sys.path:
    sys.path.insert(0, _backend_dir)
import csv
import hashlib
import json
import io
import secrets
import uuid
from datetime import datetime
from fastapi import BackgroundTasks, Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, Response, StreamingResponse
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from fastapi.staticfiles import StaticFiles
import openpyxl
from openpyxl.utils import get_column_letter
from prometheus_client import CONTENT_TYPE_LATEST, Counter, Histogram, generate_latest
from sqlalchemy import func, or_
from sqlalchemy.orm import Session
from app.config import settings
from app.database import get_db, init_db
from app.master_data import FORM_OPTIONS, KHP_MASTER_OPTIONS
from app.models import (
    Document,
    ExtractionJob,
    ExtractedField,
    KHPMasterResolution,
    KHPKelompokKegiatan,
    KHPKegiatan1,
    KHPTingkat,
    KHPJabatanPrestasi,
    KHPKegiatan2,
    KHPMasterRule,
)
from app.schemas import (
    CreateKHPRuleRequest,
    ExtractionResult,
    FieldResult,
    KHPRuleResponse,
    KHPRulesListResponse,
    OptionsResponse,
    PublicExtractionResult,
    UpdateKHPRuleRequest,
    UploadResponse,
)
from app.services.job_processor import cleanup_expired_jobs_and_uploads, process_document_job
from app.services.temporary_upload_store import upload_store
from app.services.security_guard import SecurityValidationError, inspect_and_guard_upload
from app.services.rate_limiter import enforce_upload_rate_limit

app = FastAPI(title="Certificate Autofill Prototype", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

REQUEST_COUNT = Counter("cert_autofill_requests_total", "Total HTTP requests", ["endpoint"])
UPLOAD_COUNT = Counter("cert_autofill_uploads_total", "Total uploaded PDFs")
UPLOAD_SIZE = Histogram("cert_autofill_upload_size_bytes", "PDF upload size in bytes")
logger = logging.getLogger("certificate-main")

static_dir = Path(__file__).parent / "static"
app.mount("/static", StaticFiles(directory=static_dir), name="static")


async def _storage_cleanup_loop() -> None:
    cleanup_interval = max(1, settings.storage_cleanup_interval_seconds)
    while True:
        await asyncio.sleep(cleanup_interval)
        try:
            await asyncio.to_thread(cleanup_expired_jobs_and_uploads)
        except Exception:
            logger.exception("Periodic ephemeral-storage cleanup failed.")


@app.on_event("startup")
async def on_startup() -> None:
    init_db()
    if settings.enable_khp_master_staging:
        from app.services.aucc_catalog import warm_aucc_catalog

        warm_aucc_catalog()
    cleanup_expired_jobs_and_uploads()
    app.state.storage_cleanup_task = asyncio.create_task(_storage_cleanup_loop())


@app.on_event("shutdown")
async def on_shutdown() -> None:
    task = getattr(app.state, "storage_cleanup_task", None)
    if task is not None:
        task.cancel()
        with suppress(asyncio.CancelledError):
            await task

@app.get("/", response_class=HTMLResponse)
def index() -> str:
    REQUEST_COUNT.labels(endpoint="/").inc()
    return (static_dir / "index.html").read_text(encoding="utf-8")


@app.get("/healthz")
def healthz() -> dict[str, str]:
    REQUEST_COUNT.labels(endpoint="/healthz").inc()
    return {"status": "ok"}


@app.get("/metrics")
def metrics() -> Response:
    return Response(generate_latest(), media_type=CONTENT_TYPE_LATEST)


@app.get("/api/options", response_model=OptionsResponse)
def get_options() -> OptionsResponse:
    REQUEST_COUNT.labels(endpoint="/api/options").inc()
    options = KHP_MASTER_OPTIONS if settings.enable_khp_master_staging else FORM_OPTIONS
    return OptionsResponse(options=options)


@app.post("/api/documents", response_model=UploadResponse)
def upload_document(
    request: Request,
    background_tasks: BackgroundTasks,
    tahun_akademik: str = Form(...),
    bukti_fisik: str = Form("Sertifikat"),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> UploadResponse:
    REQUEST_COUNT.labels(endpoint="/api/documents").inc()
    enforce_upload_rate_limit(request)
    if not file.filename:
        raise HTTPException(status_code=400, detail="File wajib diunggah.")

    content = file.file.read()
    max_bytes = settings.max_upload_size_mb * 1024 * 1024
    if len(content) > max_bytes:
        raise HTTPException(status_code=413, detail=f"Ukuran file maksimal {settings.max_upload_size_mb} MB.")

    try:
        validated = inspect_and_guard_upload(content, file.filename)
        temp_key = upload_store.stage_validated(validated)
    except SecurityValidationError as exc:
        raise HTTPException(status_code=400, detail=exc.message)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    document_id = str(uuid.uuid4())
    job_id = str(uuid.uuid4())

    document = Document(
        id=document_id,
        source_system="prototype_ui",
        tahun_akademik=tahun_akademik,
        bukti_fisik=bukti_fisik,
        original_file_name=file.filename,
        mime_type=validated.content_type,
        file_size=validated.cleaned_size,
        checksum_sha256=validated.checksum_sha256,
        status="queued",
    )
    job = ExtractionJob(
        id=job_id,
        document_id=document_id,
        temp_file_key=temp_key,
        status="queued",
        retry_count=0,
    )

    try:
        db.add(document)
        db.add(job)
        db.commit()
    except Exception:
        db.rollback()
        upload_store.delete(temp_key)
        raise
    # No RabbitMQ variant. The extraction job is processed either in the
    # FastAPI background task, synchronously, or by the optional DB-polling worker.
    if settings.processing_mode == "sync":
        process_document_job(job_id=job_id, document_id=document_id)
        db.refresh(document)
        response_status = document.status
    elif settings.processing_mode == "db_worker":
        response_status = "queued"
    else:
        background_tasks.add_task(process_document_job, job_id, document_id)
        response_status = "queued"

    UPLOAD_COUNT.inc()
    UPLOAD_SIZE.observe(len(content))
    return UploadResponse(document_id=document_id, job_id=job_id, status=response_status)


@app.get("/api/documents/{document_id}/result", response_model=PublicExtractionResult)
def get_result(document_id: str, db: Session = Depends(get_db)) -> PublicExtractionResult:
    REQUEST_COUNT.labels(endpoint="/api/documents/{document_id}/result").inc()
    document = db.get(Document, document_id)
    if not document:
        raise HTTPException(status_code=404, detail="Dokumen tidak ditemukan.")

    fields = db.query(ExtractedField).filter(ExtractedField.document_id == document_id).all()

    field_dict = {
        f.form_field_name: FieldResult(
            value=f.mapped_value or f.extracted_value,
            confidence=float(f.confidence),
            source=f.source,
            needs_review=bool(f.needs_review),
        )
        for f in fields
    }
    master_needs_review = False
    if settings.enable_khp_master_staging:
        resolution_row = (
            db.query(KHPMasterResolution)
            .filter(KHPMasterResolution.document_id == document_id)
            .one_or_none()
        )
        if resolution_row is not None:
            try:
                master_res = json.loads(resolution_row.resolution_json)
                master_needs_review = bool(master_res and master_res.get("status") != "resolved")
            except json.JSONDecodeError:
                master_needs_review = True

    needs_review = (
        any(item.needs_review for item in field_dict.values())
        or document.status == "needs_review"
        or master_needs_review
    )
    return PublicExtractionResult(
        document_id=document_id,
        status=document.status,
        needs_review=needs_review,
        fields=field_dict,
    )

security = HTTPBasic()


def verify_admin(credentials: HTTPBasicCredentials = Depends(security)) -> str:
    correct_username = secrets.compare_digest(credentials.username, settings.admin_username)
    correct_password = secrets.compare_digest(credentials.password, settings.admin_password)
    if not (correct_username and correct_password):
        raise HTTPException(
            status_code=401,
            detail="Akses ditolak: Kredensial admin tidak valid.",
            headers={"WWW-Authenticate": "Basic"},
        )
    return credentials.username


def _build_streaming_export_response(
    headers: list[tuple[str, str]],
    rows: list[dict],
    filename_prefix: str,
    export_format: str,
    sheet_title: str = "Data",
) -> StreamingResponse:
    header_labels = [h[0] for h in headers]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{filename_prefix}_{timestamp}"

    if export_format.lower() == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(header_labels)
        for r in rows:
            writer.writerow([r.get(h[1], "") for h in headers])
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode("utf-8")),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{filename}.csv"'},
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append(header_labels)

    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)

    for r in rows:
        ws.append([r.get(h[1], "") for h in headers])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'},
    )


@app.get("/api/admin/documents/export")
def export_documents(
    format: str = "xlsx",
    admin: str = Depends(verify_admin),
    db: Session = Depends(get_db),
):
    REQUEST_COUNT.labels(endpoint="/api/admin/documents/export").inc()
    docs = db.query(Document).order_by(Document.created_at.desc()).all()

    doc_ids = [d.id for d in docs]
    all_fields = db.query(ExtractedField).filter(ExtractedField.document_id.in_(doc_ids)).all() if doc_ids else []
    fields_by_doc: dict[str, dict[str, ExtractedField]] = {}
    for f in all_fields:
        fields_by_doc.setdefault(str(f.document_id), {})[f.form_field_name] = f

    headers = [
        ("No", "no"),
        ("ID Dokumen", "id"),
        ("Waktu Upload", "created_at"),
        ("Nama File Asli", "original_file_name"),
        ("Tahun Akademik", "tahun_akademik"),
        ("Bukti Fisik", "bukti_fisik"),
        ("Status", "status"),
        ("Engine Parser", "parser_engine"),
        ("Nama Kegiatan", "nama_kegiatan_sertifikasi"),
        ("Nomor Bukti Fisik / Sertifikat", "nomor_bukti_fisik_nomor_sertifikasi"),
        ("Penyelenggara Kegiatan", "penyelenggara_kegiatan"),
        ("Jenis Penyelenggara", "jenis_penyelenggara"),
        ("Tingkat", "tingkat"),
        ("Prestasi / Partisipasi / Jabatan", "prestasi_partisipasi_jabatan"),
        ("Kelompok Kegiatan", "kelompok_kegiatan"),
        ("Jenis Kegiatan", "jenis_kegiatan"),
        ("Waktu Mulai", "waktu_mulai_pelaksanaan"),
        ("Waktu Selesai", "waktu_selesai_pelaksanaan"),
        ("Avg Confidence", "avg_confidence"),
        ("Perlu Review", "needs_review"),
    ]

    rows = []
    for idx, doc in enumerate(docs, start=1):
        fmap = fields_by_doc.get(str(doc.id), {})

        def _val(fname: str) -> str:
            field = fmap.get(fname)
            if not field:
                return ""
            return field.mapped_value or field.extracted_value or ""

        conf_list = [float(f.confidence) for f in fmap.values() if f.confidence is not None]
        avg_conf = round(sum(conf_list) / len(conf_list), 4) if conf_list else 0.0
        has_review = any(bool(f.needs_review) for f in fmap.values()) or doc.status == "needs_review"

        rows.append({
            "no": idx,
            "id": str(doc.id),
            "created_at": doc.created_at.strftime("%Y-%m-%d %H:%M:%S") if doc.created_at else "",
            "original_file_name": doc.original_file_name,
            "tahun_akademik": doc.tahun_akademik,
            "bukti_fisik": doc.bukti_fisik,
            "status": doc.status,
            "parser_engine": doc.parser_engine or "-",
            "nama_kegiatan_sertifikasi": _val("nama_kegiatan_sertifikasi"),
            "nomor_bukti_fisik_nomor_sertifikasi": _val("nomor_bukti_fisik_nomor_sertifikasi"),
            "penyelenggara_kegiatan": _val("penyelenggara_kegiatan"),
            "jenis_penyelenggara": _val("jenis_penyelenggara"),
            "tingkat": _val("tingkat"),
            "prestasi_partisipasi_jabatan": _val("prestasi_partisipasi_jabatan"),
            "kelompok_kegiatan": _val("kelompok_kegiatan"),
            "jenis_kegiatan": _val("jenis_kegiatan"),
            "waktu_mulai_pelaksanaan": _val("waktu_mulai_pelaksanaan"),
            "waktu_selesai_pelaksanaan": _val("waktu_selesai_pelaksanaan"),
            "avg_confidence": avg_conf,
            "needs_review": "Ya" if has_review else "Tidak",
        })

    return _build_streaming_export_response(
        headers=headers,
        rows=rows,
        filename_prefix="riwayat_sertifikat",
        export_format=format,
        sheet_title="Riwayat Sertifikat",
    )


@app.get("/api/admin/khp/export")
def export_khp_master(
    table: str = "rules",
    format: str = "xlsx",
    admin: str = Depends(verify_admin),
    db: Session = Depends(get_db),
):
    REQUEST_COUNT.labels(endpoint="/api/admin/khp/export").inc()
    table_lower = table.lower().strip()

    if table_lower == "rules":
        rules = db.query(KHPMasterRule).order_by(KHPMasterRule.id.asc()).all()
        headers = [
            ("ID Rule", "id"),
            ("No Sumber", "source_no"),
            ("ID Kelompok", "id_kelompok_kegiatan"),
            ("Nama Kelompok Kegiatan", "nama_kelompok"),
            ("ID Kegiatan 1", "id_kegiatan_1"),
            ("Nama Kegiatan 1", "nama_kegiatan_1"),
            ("ID Tingkat", "id_tingkat"),
            ("Nama Tingkat", "nama_tingkat"),
            ("ID Jabatan/Prestasi", "id_jabatan_prestasi"),
            ("Nama Jabatan / Prestasi", "nama_jabatan"),
            ("Dasar Penilaian / Bukti Fisik", "dasar_penilaian"),
            ("ID Kegiatan 2", "id_kegiatan_2"),
            ("Status Aktif", "is_active"),
        ]
        rows = [
            {
                "id": r.id,
                "source_no": r.source_no,
                "id_kelompok_kegiatan": r.id_kelompok_kegiatan,
                "nama_kelompok": r.kelompok.nm_kelompok_kegiatan if r.kelompok else "",
                "id_kegiatan_1": r.id_kegiatan_1,
                "nama_kegiatan_1": r.kegiatan_1.nm_kegiatan_1 if r.kegiatan_1 else "",
                "id_tingkat": r.id_tingkat or "",
                "nama_tingkat": r.tingkat.nm_tingkat if r.tingkat else "-",
                "id_jabatan_prestasi": r.id_jabatan_prestasi or "",
                "nama_jabatan": r.jabatan_prestasi.nm_jabatan_prestasi if r.jabatan_prestasi else "-",
                "dasar_penilaian": r.dasar_penilaian,
                "id_kegiatan_2": r.id_kegiatan_2,
                "is_active": "Ya" if r.is_active else "Tidak",
            }
            for r in rules
        ]
        return _build_streaming_export_response(
            headers=headers,
            rows=rows,
            filename_prefix="khp_master_rules",
            export_format=format,
            sheet_title="Rules Penilaian KHP",
        )

    elif table_lower == "kegiatan_2":
        k2_items = db.query(KHPKegiatan2).order_by(KHPKegiatan2.id_kegiatan_2.asc()).all()
        headers = [
            ("ID Kegiatan 2", "id_kegiatan_2"),
            ("ID Kegiatan 1", "id_kegiatan_1"),
            ("Nama Kegiatan 1", "nama_kegiatan_1"),
            ("ID Tingkat", "id_tingkat"),
            ("Nama Tingkat", "nama_tingkat"),
            ("ID Jabatan/Prestasi", "id_jabatan_prestasi"),
            ("Nama Jabatan / Prestasi", "nama_jabatan"),
        ]
        rows = [
            {
                "id_kegiatan_2": k.id_kegiatan_2,
                "id_kegiatan_1": k.id_kegiatan_1,
                "nama_kegiatan_1": k.kegiatan_1.nm_kegiatan_1 if k.kegiatan_1 else "",
                "id_tingkat": k.id_tingkat or "",
                "nama_tingkat": k.tingkat.nm_tingkat if k.tingkat else "-",
                "id_jabatan_prestasi": k.id_jabatan_prestasi or "",
                "nama_jabatan": k.jabatan_prestasi.nm_jabatan_prestasi if k.jabatan_prestasi else "-",
            }
            for k in k2_items
        ]
        return _build_streaming_export_response(
            headers=headers,
            rows=rows,
            filename_prefix="khp_kegiatan_2",
            export_format=format,
            sheet_title="Master Kegiatan 2",
        )

    elif table_lower == "kegiatan_1":
        k1_items = db.query(KHPKegiatan1).order_by(KHPKegiatan1.id_kegiatan_1.asc()).all()
        headers = [
            ("ID Kegiatan 1", "id_kegiatan_1"),
            ("Nama Kegiatan 1", "nama_kegiatan_1"),
            ("ID Kelompok Kegiatan", "id_kelompok_kegiatan"),
            ("Nama Kelompok Kegiatan", "nama_kelompok"),
            ("Status Aktif", "is_aktif"),
        ]
        rows = [
            {
                "id_kegiatan_1": k.id_kegiatan_1,
                "nama_kegiatan_1": k.nm_kegiatan_1,
                "id_kelompok_kegiatan": k.id_kelompok_kegiatan,
                "nama_kelompok": k.kelompok.nm_kelompok_kegiatan if k.kelompok else "",
                "is_aktif": "Ya" if k.is_aktif else "Tidak",
            }
            for k in k1_items
        ]
        return _build_streaming_export_response(
            headers=headers,
            rows=rows,
            filename_prefix="khp_kegiatan_1",
            export_format=format,
            sheet_title="Master Kegiatan 1",
        )

    elif table_lower in {"kelompok", "kelompok_kegiatan"}:
        items = db.query(KHPKelompokKegiatan).order_by(KHPKelompokKegiatan.id_kelompok_kegiatan.asc()).all()
        headers = [
            ("ID Kelompok Kegiatan", "id"),
            ("Nama Kelompok Kegiatan", "nama"),
            ("Status Aktif", "is_aktif"),
        ]
        rows = [{"id": i.id_kelompok_kegiatan, "nama": i.nm_kelompok_kegiatan, "is_aktif": "Ya" if i.is_aktif else "Tidak"} for i in items]
        return _build_streaming_export_response(headers=headers, rows=rows, filename_prefix="khp_kelompok_kegiatan", export_format=format, sheet_title="Kelompok Kegiatan")

    elif table_lower == "tingkat":
        items = db.query(KHPTingkat).order_by(KHPTingkat.id_tingkat.asc()).all()
        headers = [("ID Tingkat", "id"), ("Nama Tingkat", "nama")]
        rows = [{"id": i.id_tingkat, "nama": i.nm_tingkat} for i in items]
        return _build_streaming_export_response(headers=headers, rows=rows, filename_prefix="khp_tingkat", export_format=format, sheet_title="Tingkat")

    elif table_lower in {"jabatan", "jabatan_prestasi"}:
        items = db.query(KHPJabatanPrestasi).order_by(KHPJabatanPrestasi.id_jabatan_prestasi.asc()).all()
        headers = [("ID Jabatan/Prestasi", "id"), ("Nama Jabatan / Prestasi", "nama")]
        rows = [{"id": i.id_jabatan_prestasi, "nama": i.nm_jabatan_prestasi} for i in items]
        return _build_streaming_export_response(headers=headers, rows=rows, filename_prefix="khp_jabatan_prestasi", export_format=format, sheet_title="Jabatan Prestasi")

    else:
        raise HTTPException(
            status_code=400,
            detail=f"Tabel '{table}' tidak didukung. Pilihan: rules, kegiatan_2, kegiatan_1, kelompok, tingkat, jabatan.",
        )


@app.get("/api/admin/khp/rules", response_model=KHPRulesListResponse)
def list_khp_rules(
    search: str | None = None,
    id_kelompok_kegiatan: int | None = None,
    id_kegiatan_1: int | None = None,
    id_tingkat: int | None = None,
    is_active: bool | None = None,
    limit: int = 100,
    offset: int = 0,
    admin: str = Depends(verify_admin),
    db: Session = Depends(get_db),
):
    REQUEST_COUNT.labels(endpoint="/api/admin/khp/rules").inc()
    query = db.query(KHPMasterRule)
    if id_kelompok_kegiatan is not None:
        query = query.filter(KHPMasterRule.id_kelompok_kegiatan == id_kelompok_kegiatan)
    if id_kegiatan_1 is not None:
        query = query.filter(KHPMasterRule.id_kegiatan_1 == id_kegiatan_1)
    if id_tingkat is not None:
        query = query.filter(KHPMasterRule.id_tingkat == id_tingkat)
    if is_active is not None:
        query = query.filter(KHPMasterRule.is_active == is_active)
    if search:
        s = f"%{search}%"
        query = query.join(KHPMasterRule.kegiatan_1).filter(
            or_(
                KHPKegiatan1.nm_kegiatan_1.ilike(s),
                KHPMasterRule.dasar_penilaian.ilike(s),
            )
        )
    total = query.count()
    rules = query.order_by(KHPMasterRule.id.asc()).offset(offset).limit(limit).all()
    items = [
        KHPRuleResponse(
            id=r.id,
            source_no=r.source_no,
            id_kelompok_kegiatan=r.id_kelompok_kegiatan,
            nama_kelompok_kegiatan=r.kelompok.nm_kelompok_kegiatan if r.kelompok else None,
            id_kegiatan_1=r.id_kegiatan_1,
            nama_kegiatan_1=r.kegiatan_1.nm_kegiatan_1 if r.kegiatan_1 else None,
            id_tingkat=r.id_tingkat,
            nama_tingkat=r.tingkat.nm_tingkat if r.tingkat else None,
            id_jabatan_prestasi=r.id_jabatan_prestasi,
            nama_jabatan_prestasi=r.jabatan_prestasi.nm_jabatan_prestasi if r.jabatan_prestasi else None,
            dasar_penilaian=r.dasar_penilaian,
            id_kegiatan_2=r.id_kegiatan_2,
            is_active=r.is_active,
            created_at=r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else None,
            updated_at=r.updated_at.strftime("%Y-%m-%d %H:%M:%S") if r.updated_at else None,
        )
        for r in rules
    ]
    return KHPRulesListResponse(total=total, items=items)


@app.post("/api/admin/khp/rules", response_model=KHPRuleResponse, status_code=201)
def create_khp_rule(
    req: CreateKHPRuleRequest,
    admin: str = Depends(verify_admin),
    db: Session = Depends(get_db),
):
    REQUEST_COUNT.labels(endpoint="/api/admin/khp/rules").inc()
    # Validasi keberadaan dimensi
    kelompok = db.query(KHPKelompokKegiatan).filter_by(id_kelompok_kegiatan=req.id_kelompok_kegiatan).first()
    if not kelompok:
        raise HTTPException(status_code=400, detail=f"Kelompok kegiatan ID {req.id_kelompok_kegiatan} tidak ditemukan.")
    keg1 = db.query(KHPKegiatan1).filter_by(id_kegiatan_1=req.id_kegiatan_1).first()
    if not keg1:
        raise HTTPException(status_code=400, detail=f"Kegiatan 1 ID {req.id_kegiatan_1} tidak ditemukan.")

    if req.id_tingkat is not None:
        if not db.query(KHPTingkat).filter_by(id_tingkat=req.id_tingkat).first():
            raise HTTPException(status_code=400, detail=f"Tingkat ID {req.id_tingkat} tidak ditemukan.")
    if req.id_jabatan_prestasi is not None:
        if not db.query(KHPJabatanPrestasi).filter_by(id_jabatan_prestasi=req.id_jabatan_prestasi).first():
            raise HTTPException(status_code=400, detail=f"Jabatan prestasi ID {req.id_jabatan_prestasi} tidak ditemukan.")

    # Cek duplikasi rule persis
    existing_rule = db.query(KHPMasterRule).filter_by(
        id_kelompok_kegiatan=req.id_kelompok_kegiatan,
        id_kegiatan_1=req.id_kegiatan_1,
        id_tingkat=req.id_tingkat,
        id_jabatan_prestasi=req.id_jabatan_prestasi,
    ).first()
    if existing_rule:
        raise HTTPException(
            status_code=409,
            detail=f"Rule untuk kombinasi ini sudah ada (Rule ID {existing_rule.id}). Gunakan update jika ingin mengubah.",
        )

    # Resolve atau generate id_kegiatan_2
    id_kegiatan_2 = req.id_kegiatan_2
    if id_kegiatan_2 is None:
        k2 = db.query(KHPKegiatan2).filter_by(
            id_kegiatan_1=req.id_kegiatan_1,
            id_tingkat=req.id_tingkat,
            id_jabatan_prestasi=req.id_jabatan_prestasi,
        ).first()
        if k2:
            id_kegiatan_2 = k2.id_kegiatan_2
        else:
            max_k2 = db.query(func.max(KHPKegiatan2.id_kegiatan_2)).scalar() or 0
            id_kegiatan_2 = max_k2 + 1
            new_k2 = KHPKegiatan2(
                id_kegiatan_2=id_kegiatan_2,
                id_kegiatan_1=req.id_kegiatan_1,
                id_tingkat=req.id_tingkat,
                id_jabatan_prestasi=req.id_jabatan_prestasi,
            )
            db.add(new_k2)
            db.flush()

    max_source_no = db.query(func.max(KHPMasterRule.source_no)).scalar() or 0
    new_rule = KHPMasterRule(
        source_no=max_source_no + 1,
        id_kelompok_kegiatan=req.id_kelompok_kegiatan,
        id_kegiatan_1=req.id_kegiatan_1,
        id_tingkat=req.id_tingkat,
        id_jabatan_prestasi=req.id_jabatan_prestasi,
        dasar_penilaian=req.dasar_penilaian.strip(),
        id_kegiatan_2=id_kegiatan_2,
        is_active=req.is_active,
    )
    db.add(new_rule)
    db.commit()
    db.refresh(new_rule)

    return KHPRuleResponse(
        id=new_rule.id,
        source_no=new_rule.source_no,
        id_kelompok_kegiatan=new_rule.id_kelompok_kegiatan,
        nama_kelompok_kegiatan=new_rule.kelompok.nm_kelompok_kegiatan if new_rule.kelompok else None,
        id_kegiatan_1=new_rule.id_kegiatan_1,
        nama_kegiatan_1=new_rule.kegiatan_1.nm_kegiatan_1 if new_rule.kegiatan_1 else None,
        id_tingkat=new_rule.id_tingkat,
        nama_tingkat=new_rule.tingkat.nm_tingkat if new_rule.tingkat else None,
        id_jabatan_prestasi=new_rule.id_jabatan_prestasi,
        nama_jabatan_prestasi=new_rule.jabatan_prestasi.nm_jabatan_prestasi if new_rule.jabatan_prestasi else None,
        dasar_penilaian=new_rule.dasar_penilaian,
        id_kegiatan_2=new_rule.id_kegiatan_2,
        is_active=new_rule.is_active,
        created_at=new_rule.created_at.strftime("%Y-%m-%d %H:%M:%S") if new_rule.created_at else None,
        updated_at=new_rule.updated_at.strftime("%Y-%m-%d %H:%M:%S") if new_rule.updated_at else None,
    )


@app.patch("/api/admin/khp/rules/{rule_id}", response_model=KHPRuleResponse)
def update_khp_rule(
    rule_id: int,
    req: UpdateKHPRuleRequest,
    admin: str = Depends(verify_admin),
    db: Session = Depends(get_db),
):
    REQUEST_COUNT.labels(endpoint="/api/admin/khp/rules/{rule_id}").inc()
    rule = db.query(KHPMasterRule).filter_by(id=rule_id).first()
    if not rule:
        raise HTTPException(status_code=404, detail=f"Rule ID {rule_id} tidak ditemukan.")
    if req.dasar_penilaian is not None:
        rule.dasar_penilaian = req.dasar_penilaian.strip()
    if req.is_active is not None:
        rule.is_active = req.is_active
    db.commit()
    db.refresh(rule)
    return KHPRuleResponse(
        id=rule.id,
        source_no=rule.source_no,
        id_kelompok_kegiatan=rule.id_kelompok_kegiatan,
        nama_kelompok_kegiatan=rule.kelompok.nm_kelompok_kegiatan if rule.kelompok else None,
        id_kegiatan_1=rule.id_kegiatan_1,
        nama_kegiatan_1=rule.kegiatan_1.nm_kegiatan_1 if rule.kegiatan_1 else None,
        id_tingkat=rule.id_tingkat,
        nama_tingkat=rule.tingkat.nm_tingkat if rule.tingkat else None,
        id_jabatan_prestasi=rule.id_jabatan_prestasi,
        nama_jabatan_prestasi=rule.jabatan_prestasi.nm_jabatan_prestasi if rule.jabatan_prestasi else None,
        dasar_penilaian=rule.dasar_penilaian,
        id_kegiatan_2=rule.id_kegiatan_2,
        is_active=rule.is_active,
        created_at=rule.created_at.strftime("%Y-%m-%d %H:%M:%S") if rule.created_at else None,
        updated_at=rule.updated_at.strftime("%Y-%m-%d %H:%M:%S") if rule.updated_at else None,
    )


@app.delete("/api/admin/khp/rules/{rule_id}")
def delete_khp_rule(
    rule_id: int,
    hard: bool = False,
    admin: str = Depends(verify_admin),
    db: Session = Depends(get_db),
):
    REQUEST_COUNT.labels(endpoint="/api/admin/khp/rules/{rule_id}").inc()
    rule = db.query(KHPMasterRule).filter_by(id=rule_id).first()
    if not rule:
        raise HTTPException(status_code=404, detail=f"Rule ID {rule_id} tidak ditemukan.")
    if hard:
        db.delete(rule)
        db.commit()
        return {"message": f"Rule ID {rule_id} berhasil dihapus permanen."}
    rule.is_active = False
    db.commit()
    return {"message": f"Rule ID {rule_id} berhasil dinonaktifkan (soft delete)."}
