"""Unit tests for /api/admin/documents/export endpoint with HTTP Basic Auth.

Tests:
1. Rejection of unauthenticated requests (401 Unauthorized).
2. Rejection of requests with invalid credentials (401 Unauthorized).
3. Successful XLSX export with valid credentials and verify workbook structure.
4. Successful CSV export with valid credentials and verify CSV content.
"""

import io
import csv
import uuid
from datetime import datetime, timezone
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.pool import StaticPool
from sqlalchemy.orm import sessionmaker

from app.database import Base, get_db
from app.main import app
from app.models import Document, ExtractedField


@pytest.fixture
def db_session():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()
@pytest.fixture
def client(db_session, monkeypatch):
    monkeypatch.setattr("app.main.init_db", lambda: None)
    monkeypatch.setattr("app.main.cleanup_expired_jobs_and_uploads", lambda: None)

    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def test_export_unauthorized_without_auth(client):
    response = client.get("/api/admin/documents/export")
    assert response.status_code == 401
    assert "WWW-Authenticate" in response.headers


def test_export_unauthorized_with_wrong_password(client):
    response = client.get("/api/admin/documents/export", auth=("admin", "wrong_password"))
    assert response.status_code == 401


def test_export_unauthorized_with_wrong_username(client):
    response = client.get("/api/admin/documents/export", auth=("hacker", "admin123"))
    assert response.status_code == 401


def test_export_xlsx_success(client, db_session):
    doc_id = str(uuid.uuid4())
    doc = Document(
        id=doc_id,
        source_system="manual_test",
        tahun_akademik="2024/2025 - Ganjil",
        bukti_fisik="Sertifikat",
        original_file_name="sertifikat_lomba_ai.pdf",
        mime_type="application/pdf",
        file_size=102400,
        checksum_sha256="dummy_sha256",
        status="completed",
        parser_engine="gemini-prod",
        created_at=datetime(2026, 3, 10, 10, 30, 0, tzinfo=timezone.utc),
    )
    db_session.add(doc)

    f1 = ExtractedField(
        id=str(uuid.uuid4()),
        document_id=doc_id,
        form_field_name="nama_kegiatan_sertifikasi",
        extracted_value="Lomba AI Mahasiswa 2026",
        mapped_value="Lomba AI Mahasiswa 2026",
        confidence=0.96,
        source="gemini",
        needs_review=False,
    )
    f2 = ExtractedField(
        id=str(uuid.uuid4()),
        document_id=doc_id,
        form_field_name="penyelenggara_kegiatan",
        extracted_value="BEM Fakultas Teknologi Maju dan Multidisiplin",
        mapped_value="BEM Fakultas Teknologi Maju dan Multidisiplin",
        confidence=0.92,
        source="gemini",
        needs_review=False,
    )
    f3 = ExtractedField(
        id=str(uuid.uuid4()),
        document_id=doc_id,
        form_field_name="tingkat",
        extracted_value="Nasional",
        mapped_value="Nasional",
        confidence=0.98,
        source="gemini",
        needs_review=False,
    )
    db_session.add_all([f1, f2, f3])
    db_session.commit()

    response = client.get("/api/admin/documents/export?format=xlsx", auth=("admin", "admin123"))
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers.get("content-type", "")
    assert "attachment; filename=\"riwayat_sertifikat_" in response.headers.get("content-disposition", "")

    # Inspect the returned XLSX workbook
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Riwayat Sertifikat" in wb.sheetnames
    ws = wb["Riwayat Sertifikat"]

    # Header check
    headers = [cell.value for cell in ws[1]]
    assert "Nama File Asli" in headers
    assert "Nama Kegiatan" in headers
    assert "Penyelenggara Kegiatan" in headers
    assert "Tingkat" in headers

    # Row data check
    row_values = [cell.value for cell in ws[2]]
    assert "sertifikat_lomba_ai.pdf" in row_values
    assert "Lomba AI Mahasiswa 2026" in row_values
    assert "BEM Fakultas Teknologi Maju dan Multidisiplin" in row_values
    assert "Nasional" in row_values


def test_export_csv_success(client, db_session):
    doc_id = str(uuid.uuid4())
    doc = Document(
        id=doc_id,
        source_system="manual_test",
        tahun_akademik="2024/2025 - Genap",
        bukti_fisik="Sertifikat",
        original_file_name="cert_webinar.pdf",
        mime_type="application/pdf",
        file_size=51200,
        checksum_sha256="dummy_sha256_2",
        status="needs_review",
        parser_engine="rules-v4.2",
        created_at=datetime(2026, 3, 10, 11, 0, 0, tzinfo=timezone.utc),
    )
    db_session.add(doc)

    f = ExtractedField(
        id=str(uuid.uuid4()),
        document_id=doc_id,
        form_field_name="nama_kegiatan_sertifikasi",
        extracted_value="Webinar Teknologi Cloud",
        mapped_value="Webinar Teknologi Cloud",
        confidence=0.75,
        source="rules",
        needs_review=True,
    )
    db_session.add(f)
    db_session.commit()

    response = client.get("/api/admin/documents/export?format=csv", auth=("admin", "admin123"))
    assert response.status_code == 200
    assert "text/csv" in response.headers.get("content-type", "")

    content = response.content.decode("utf-8")
    reader = list(csv.reader(io.StringIO(content)))
    assert len(reader) >= 2
    header = reader[0]
    row = reader[1]

    name_idx = header.index("Nama File Asli")
    act_idx = header.index("Nama Kegiatan")
    review_idx = header.index("Perlu Review")

    assert row[name_idx] == "cert_webinar.pdf"
    assert row[act_idx] == "Webinar Teknologi Cloud"
    assert row[review_idx] == "Ya"
