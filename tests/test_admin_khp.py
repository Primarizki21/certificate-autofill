"""Unit tests for Admin KHP Master export and rule management endpoints.

Tests:
1. Exporting KHP tables (rules, kegiatan_2, kegiatan_1, kelompok, tingkat, jabatan) in XLSX & CSV.
2. Handling invalid table parameter (400 Bad Request).
3. Enforcing authentication (401 Unauthorized).
4. Listing rules with search and filters.
5. Creating new rules with automatic kegiatan_2 resolution.
6. Preventing duplicate rules (409 Conflict).
7. Validating foreign key existence (400 Bad Request).
8. Updating rule attributes (PATCH).
9. Soft-deleting and hard-deleting rules (DELETE).
"""

import csv
import io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app
from app.models import (
    KHPKelompokKegiatan,
    KHPKegiatan1,
    KHPTingkat,
    KHPJabatanPrestasi,
    KHPKegiatan2,
    KHPMasterRule,
)


@pytest.fixture
def db_session():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    session = TestingSessionLocal()

    # Seed initial test data
    kel1 = KHPKelompokKegiatan(id_kelompok_kegiatan=1, nm_kelompok_kegiatan="Wajib Universitas", is_aktif=True)
    kel2 = KHPKelompokKegiatan(id_kelompok_kegiatan=2, nm_kelompok_kegiatan="Organisasi dan Kepemimpinan", is_aktif=True)
    session.add_all([kel1, kel2])

    keg1 = KHPKegiatan1(id_kegiatan_1=41, nm_kegiatan_1="PKKMB", id_kelompok_kegiatan=1, is_aktif=True)
    keg2 = KHPKegiatan1(id_kegiatan_1=67, nm_kegiatan_1="Pengurus Organisasi", id_kelompok_kegiatan=2, is_aktif=True)
    session.add_all([keg1, keg2])

    t1 = KHPTingkat(id_tingkat=1, nm_tingkat="Internasional")
    t2 = KHPTingkat(id_tingkat=2, nm_tingkat="Nasional")
    session.add_all([t1, t2])

    j1 = KHPJabatanPrestasi(id_jabatan_prestasi=1, nm_jabatan_prestasi="Ketua")
    j2 = KHPJabatanPrestasi(id_jabatan_prestasi=2, nm_jabatan_prestasi="Wakil Ketua")
    session.add_all([j1, j2])

    k2_1 = KHPKegiatan2(id_kegiatan_2=100, id_kegiatan_1=67, id_tingkat=2, id_jabatan_prestasi=1)
    k2_2 = KHPKegiatan2(id_kegiatan_2=101, id_kegiatan_1=67, id_tingkat=2, id_jabatan_prestasi=2)
    session.add_all([k2_1, k2_2])

    rule1 = KHPMasterRule(
        id=1,
        source_no=10,
        id_kelompok_kegiatan=2,
        id_kegiatan_1=67,
        id_tingkat=2,
        id_jabatan_prestasi=1,
        dasar_penilaian="SK / Sertifikat",
        id_kegiatan_2=100,
        is_active=True,
    )
    session.add(rule1)
    session.commit()

    try:
        yield session
    finally:
        session.close()


@pytest.fixture
def client(db_session, monkeypatch):
    monkeypatch.setattr("app.main.init_db", lambda: None)
    monkeypatch.setattr("app.main.cleanup_expired_jobs_and_uploads", lambda: None)

    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


AUTH = ("admin", "admin123")


def test_khp_export_unauthorized(client):
    res = client.get("/api/admin/khp/export")
    assert res.status_code == 401


def test_khp_export_invalid_table(client):
    res = client.get("/api/admin/khp/export?table=invalid_table", auth=AUTH)
    assert res.status_code == 400
    assert "tidak didukung" in res.json()["detail"]


def test_khp_export_rules_xlsx(client):
    res = client.get("/api/admin/khp/export?table=rules&format=xlsx", auth=AUTH)
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    assert "khp_master_rules_" in res.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Rules Penilaian KHP" in wb.sheetnames
    ws = wb["Rules Penilaian KHP"]

    headers = [c.value for c in ws[1]]
    assert "ID Rule" in headers
    assert "Nama Kegiatan 1" in headers
    assert "Nama Tingkat" in headers
    assert "Dasar Penilaian / Bukti Fisik" in headers

    row = [c.value for c in ws[2]]
    assert 1 in row
    assert "Pengurus Organisasi" in row
    assert "Nasional" in row
    assert "Ketua" in row


def test_khp_export_rules_csv(client):
    res = client.get("/api/admin/khp/export?table=rules&format=csv", auth=AUTH)
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]

    content = res.content.decode("utf-8")
    lines = list(csv.reader(io.StringIO(content)))
    assert len(lines) >= 2
    header = lines[0]
    row = lines[1]
    assert "ID Rule" in header
    assert row[header.index("Nama Kegiatan 1")] == "Pengurus Organisasi"
    assert row[header.index("Status Aktif")] == "Ya"


def test_khp_export_other_tables_csv(client):
    for tbl in ["kegiatan_2", "kegiatan_1", "kelompok", "tingkat", "jabatan"]:
        res = client.get(f"/api/admin/khp/export?table={tbl}&format=csv", auth=AUTH)
        assert res.status_code == 200
        assert "text/csv" in res.headers["content-type"]
        lines = list(csv.reader(io.StringIO(res.content.decode("utf-8"))))
        assert len(lines) >= 2


def test_khp_rules_list_and_search(client):
    res = client.get("/api/admin/khp/rules", auth=AUTH)
    assert res.status_code == 200
    data = res.json()
    assert data["total"] == 1
    assert data["items"][0]["nama_kegiatan_1"] == "Pengurus Organisasi"

    # Search match
    res_search = client.get("/api/admin/khp/rules?search=Pengurus", auth=AUTH)
    assert res_search.status_code == 200
    assert res_search.json()["total"] == 1

    # Search no match
    res_no = client.get("/api/admin/khp/rules?search=NonExistentWord", auth=AUTH)
    assert res_no.status_code == 200
    assert res_no.json()["total"] == 0


def test_khp_rules_create_success_existing_k2(client):
    payload = {
        "id_kelompok_kegiatan": 2,
        "id_kegiatan_1": 67,
        "id_tingkat": 2,
        "id_jabatan_prestasi": 2,
        "dasar_penilaian": "Sertifikat Resmi",
        "is_active": True,
    }
    res = client.post("/api/admin/khp/rules", json=payload, auth=AUTH)
    assert res.status_code == 201
    created = res.json()
    assert created["id_kegiatan_2"] == 101
    assert created["nama_jabatan_prestasi"] == "Wakil Ketua"
    assert created["dasar_penilaian"] == "Sertifikat Resmi"


def test_khp_rules_create_success_new_k2(client):
    payload = {
        "id_kelompok_kegiatan": 1,
        "id_kegiatan_1": 41,
        "id_tingkat": 1,
        "id_jabatan_prestasi": 1,
        "dasar_penilaian": "Sertifikat PKKMB Internasional",
        "is_active": True,
    }
    res = client.post("/api/admin/khp/rules", json=payload, auth=AUTH)
    assert res.status_code == 201
    created = res.json()
    assert created["id_kegiatan_2"] > 101  # Auto-generated
    assert created["nama_kegiatan_1"] == "PKKMB"


def test_khp_rules_create_duplicate_conflict(client):
    payload = {
        "id_kelompok_kegiatan": 2,
        "id_kegiatan_1": 67,
        "id_tingkat": 2,
        "id_jabatan_prestasi": 1,
        "dasar_penilaian": "Duplikat",
    }
    res = client.post("/api/admin/khp/rules", json=payload, auth=AUTH)
    assert res.status_code == 409
    assert "sudah ada" in res.json()["detail"]


def test_khp_rules_create_invalid_foreign_key(client):
    payload = {
        "id_kelompok_kegiatan": 9999,  # invalid
        "id_kegiatan_1": 67,
        "dasar_penilaian": "Test",
    }
    res = client.post("/api/admin/khp/rules", json=payload, auth=AUTH)
    assert res.status_code == 400
    assert "Kelompok kegiatan ID 9999 tidak ditemukan" in res.json()["detail"]


def test_khp_rules_update(client):
    res = client.patch(
        "/api/admin/khp/rules/1",
        json={"dasar_penilaian": "SK Terverifikasi", "is_active": False},
        auth=AUTH,
    )
    assert res.status_code == 200
    updated = res.json()
    assert updated["dasar_penilaian"] == "SK Terverifikasi"
    assert updated["is_active"] is False

    # 404 on non-existent
    res_404 = client.patch("/api/admin/khp/rules/99999", json={"is_active": True}, auth=AUTH)
    assert res_404.status_code == 404


def test_khp_rules_delete(client):
    # Soft delete
    res_soft = client.delete("/api/admin/khp/rules/1", auth=AUTH)
    assert res_soft.status_code == 200
    assert "dinonaktifkan" in res_soft.json()["message"]

    # Verify soft delete
    rule_res = client.get("/api/admin/khp/rules?is_active=false", auth=AUTH)
    assert rule_res.json()["total"] == 1

    # Hard delete
    res_hard = client.delete("/api/admin/khp/rules/1?hard=true", auth=AUTH)
    assert res_hard.status_code == 200
    assert "berhasil dihapus permanen" in res_hard.json()["message"]

    # Verify deleted
    assert client.get("/api/admin/khp/rules", auth=AUTH).json()["total"] == 0
